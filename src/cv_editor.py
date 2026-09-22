"""
Extraction et édition "en place" d'un CV Word (.docx).

Principe : on repère chaque bloc de texte du document (paragraphes du corps,
cellules de tableaux, en-têtes/pieds de page) et on leur donne un identifiant
stable. L'IA reçoit ces blocs, propose un nouveau texte pour ceux qui doivent
changer, et on réinjecte ce texte dans le document d'origine en gardant la
mise en forme (police, gras, taille...) du premier "run" de chaque paragraphe.

Limites connues (voir README) :
- Le texte contenu dans des zones de texte / images / SmartArt n'est pas
  détecté (python-docx ne les lit pas) : privilégie un CV avec du texte
  "normal" (paragraphes, tableaux), ce qui est de toute façon recommandé
  pour la compatibilité avec les logiciels de tri de CV (ATS).
- Si un paragraphe contient plusieurs styles différents dans la même phrase
  (ex: un mot en gras au milieu d'une phrase normale), seule la mise en
  forme du tout premier mot est conservée après remplacement.
"""

from dataclasses import dataclass
from pathlib import Path
from typing import Iterator

from docx import Document
from docx.text.paragraph import Paragraph


@dataclass
class BlocTexte:
    id: str
    texte: str
    longueur: int


def _paragraphes_non_vides(paragraphes, prefixe: str) -> Iterator[tuple[str, Paragraph]]:
    for i, p in enumerate(paragraphes):
        if p.text.strip():
            yield f"{prefixe}{i}", p


def _iterer_paragraphes_identifies(document: Document) -> Iterator[tuple[str, Paragraph]]:
    """Parcourt le document dans un ordre déterministe et stable.

    Utilisé à la fois pour l'extraction et pour la réinjection : les deux
    passes doivent absolument produire les mêmes identifiants pour les mêmes
    paragraphes, donc toute la logique de parcours vit ici, une seule fois.
    """
    # Corps du document
    yield from _paragraphes_non_vides(document.paragraphs, "p")

    # Tableaux (ex: mise en page en colonnes pour les compétences)
    for ti, table in enumerate(document.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                yield from _paragraphes_non_vides(
                    cell.paragraphs, f"t{ti}_r{ri}_c{ci}_p"
                )

    # En-têtes / pieds de page (parfois utilisés pour le nom ou un titre)
    for si, section in enumerate(document.sections):
        yield from _paragraphes_non_vides(section.header.paragraphs, f"h{si}_p")
        yield from _paragraphes_non_vides(section.footer.paragraphs, f"f{si}_p")


def extraire_blocs(chemin_docx: Path) -> list[BlocTexte]:
    """Retourne la liste des blocs de texte détectés dans le CV, dans l'ordre."""
    document = Document(str(chemin_docx))
    return [
        BlocTexte(id=identifiant, texte=p.text, longueur=len(p.text))
        for identifiant, p in _iterer_paragraphes_identifies(document)
    ]


def texte_brut(blocs: list[BlocTexte]) -> str:
    """Vue à plat du contenu du CV, pour donner du contexte global à l'IA."""
    return "\n".join(b.texte for b in blocs)


def _remplacer_texte_paragraphe(paragraphe: Paragraph, nouveau_texte: str) -> None:
    runs = paragraphe.runs
    if not runs:
        paragraphe.add_run(nouveau_texte)
        return
    runs[0].text = nouveau_texte
    for run in runs[1:]:
        run.text = ""


def appliquer_adaptations(
    chemin_source: Path, chemin_destination: Path, adaptations: dict[str, str]
) -> list[str]:
    """Crée une copie du CV avec les blocs listés dans `adaptations` remplacés.

    Le fichier source (le CV de référence) n'est jamais modifié : le résultat
    est écrit dans `chemin_destination`. Retourne la liste des identifiants de
    bloc effectivement trouvés et modifiés (pour signaler ceux qui, le cas
    échéant, n'auraient pas été retrouvés si le CV de référence a changé
    entre-temps).
    """
    document = Document(str(chemin_source))
    modifies: list[str] = []
    for identifiant, paragraphe in _iterer_paragraphes_identifies(document):
        if identifiant in adaptations:
            _remplacer_texte_paragraphe(paragraphe, adaptations[identifiant])
            modifies.append(identifiant)

    chemin_destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(chemin_destination))
    return modifies
