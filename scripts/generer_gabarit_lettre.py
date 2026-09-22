"""
Génère le gabarit Word de la lettre de motivation : templates/lettre_template.docx

Ce script construit le fichier .docx avec python-docx plutôt que de le taper
à la main dans Word. Raison : docxtpl (qui remplit le gabarit) repère les
balises {{ ... }} et {%p ... %} dans le texte, mais Word a tendance à couper
un texte tapé au clavier en plusieurs morceaux internes ("runs"), ce qui casse
les balises de façon invisible à l'œil. En générant le fichier ici, chaque
balise est écrite en un seul bloc, donc le gabarit reste fiable.

Pour changer la mise en page de la lettre (police, taille, alignement,
espacements...), modifie ce script puis relance-le :

    python scripts/generer_gabarit_lettre.py

Ne modifie jamais directement templates/lettre_template.docx dans Word :
la prochaine régénération écraserait tes changements, et une édition
manuelle des balises risquerait de les casser.
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

RACINE = Path(__file__).resolve().parent.parent
SORTIE = RACINE / "templates" / "lettre_template.docx"

POLICE = "Calibri"
TAILLE_CORPS = Pt(11)
TAILLE_PETIT = Pt(10)


def paragraphe(doc, texte, *, taille=TAILLE_CORPS, gras=False, alignement=None, espace_apres=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(espace_apres)
    if alignement is not None:
        p.alignment = alignement
    run = p.add_run(texte)
    run.font.name = POLICE
    run.font.size = taille
    run.bold = gras
    return p


def ligne_vide(doc, espace_apres=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(espace_apres)
    return p


def construire():
    doc = Document()

    # Marges et police par défaut du document
    style_normal = doc.styles["Normal"]
    style_normal.font.name = POLICE
    style_normal.font.size = TAILLE_CORPS
    for section in doc.sections:
        section.top_margin = Pt(56)
        section.bottom_margin = Pt(56)
        section.left_margin = Pt(70)
        section.right_margin = Pt(70)

    gauche = WD_ALIGN_PARAGRAPH.LEFT
    droite = WD_ALIGN_PARAGRAPH.RIGHT
    justifie = WD_ALIGN_PARAGRAPH.JUSTIFY

    # --- Bloc expéditeur (haut gauche) ---
    paragraphe(doc, "{{ expediteur_nom }}", taille=TAILLE_PETIT, gras=True, alignement=gauche, espace_apres=0)
    paragraphe(doc, "{{ expediteur_adresse }}", taille=TAILLE_PETIT, alignement=gauche, espace_apres=0)
    paragraphe(doc, "{{ expediteur_cp_ville }}", taille=TAILLE_PETIT, alignement=gauche, espace_apres=0)
    paragraphe(doc, "{{ expediteur_telephone }}", taille=TAILLE_PETIT, alignement=gauche, espace_apres=0)
    paragraphe(doc, "{{ expediteur_email }}", taille=TAILLE_PETIT, alignement=gauche, espace_apres=18)

    # --- Ville, date (droite) ---
    paragraphe(doc, "{{ ville_signature }}, le {{ date_lettre }}", taille=TAILLE_PETIT, alignement=droite, espace_apres=18)

    # --- Bloc destinataire (droite) ---
    paragraphe(doc, "{{ entreprise }}", taille=TAILLE_PETIT, gras=True, alignement=droite, espace_apres=0)
    paragraphe(doc, "{{ entreprise_adresse }}", taille=TAILLE_PETIT, alignement=droite, espace_apres=24)

    # --- Objet ---
    paragraphe(doc, "Objet : {{ objet }}", gras=True, alignement=gauche, espace_apres=18)

    # --- Formule d'appel ---
    paragraphe(doc, "{{ formule_appel }}", alignement=gauche, espace_apres=12)

    # --- Corps de la lettre : boucle paragraphe-par-paragraphe (docxtpl "{%p %}") ---
    ligne_boucle_debut = doc.add_paragraph("{%p for paragraphe_corps in paragraphes %}")
    p_corps = doc.add_paragraph()
    p_corps.paragraph_format.space_after = Pt(12)
    p_corps.paragraph_format.line_spacing = 1.15
    p_corps.alignment = justifie
    run_corps = p_corps.add_run("{{ paragraphe_corps }}")
    run_corps.font.name = POLICE
    run_corps.font.size = TAILLE_CORPS
    ligne_boucle_fin = doc.add_paragraph("{%p endfor %}")

    # --- Formule de politesse + signature ---
    paragraphe(doc, "{{ formule_politesse }}", alignement=gauche, espace_apres=36)
    paragraphe(doc, "{{ expediteur_nom }}", gras=True, alignement=gauche, espace_apres=0)

    SORTIE.parent.mkdir(parents=True, exist_ok=True)
    doc.save(SORTIE)
    print(f"Gabarit généré : {SORTIE}")


if __name__ == "__main__":
    construire()
